from django.db.models.functions import Coalesce

from .forms import PriceForm, AuthorForm, TagForm
from django.contrib.auth.decorators import (
    login_required
)
from django.db.models import Q
from django.http import JsonResponse
from django.template.loader import render_to_string
from shop.models import Category, Price
from django.db.models import Sum
from django.utils.dateparse import parse_date
from django.http import HttpResponse

from openpyxl import Workbook
from reportlab.pdfgen import canvas
from django.shortcuts import render, get_object_or_404, redirect

from orders.models import (
    Order,
    OrderItem
)
from django.utils import timezone
from datetime import timedelta
from django.db.models import Count
from django.db.models import Sum

from shop.models import (
    Product,
    ProductImage,
    Book,
    BoardGame,
    Stationery,
    Author,
    Genre,
    Tag
)

from .decorators import admin_required

from .forms import (
    ProductForm,
    BookForm,
    BoardGameForm,
    StationeryForm
)

from admin_panel.constant import (
    CATEGORY_MAP,
    RELATED_MAP
)
from shop.models import Genre
from .forms import GenreForm
from django.contrib import messages

from shop.models import Supplier
from .forms import SupplierForm

from django.db.models import Sum, Count, Q
from django.utils import timezone
from datetime import timedelta

def get_period(period):
    now = timezone.now()

    if period == 'day':
        return now - timedelta(days=1)
    if period == 'week':
        return now - timedelta(days=7)
    if period == 'month':
        return now - timedelta(days=30)
    if period == 'year':
        return now - timedelta(days=365)

    return now - timedelta(days=30)

def get_sales_by_period(days):
    since = timezone.now() - timedelta(days=days)

    return OrderItem.objects.filter(
        order__created__gte=since
    ).aggregate(
        total=Sum('quantity')
    )['total'] or 0


def supplier_list(request):
    q = request.GET.get('q')

    suppliers = Supplier.objects.filter(is_active=True)
    archived_suppliers = Supplier.objects.filter(is_active=False)

    if q:
        suppliers = suppliers.filter(
            name__icontains=q
        )

        archived_suppliers = archived_suppliers.filter(
            name__icontains=q
        )

    return render(
        request,
        'admin_panel/dictionaries/suppliers/suppliers.html',
        {
            'suppliers': suppliers,
            'archived_suppliers': archived_suppliers,
        }
    )

def supplier_create(request):
    form = SupplierForm(request.POST or None)

    if form.is_valid():
        form.save()
        return redirect('admin_panel:supplier_list')

    return render(request, 'admin_panel/dictionaries/suppliers/supplier_form.html', {'form': form})

def supplier_delete(request, id):
    supplier = get_object_or_404(Supplier, id=id)

    used = (
        Book.objects.filter(publisher=supplier).exists() or
        BoardGame.objects.filter(publisher=supplier).exists() or
        Stationery.objects.filter(supplier=supplier).exists()
    )

    if used:
        supplier.is_active = False
        supplier.save()
    else:
        supplier.delete()

    return redirect('admin_panel:supplier_list')

def supplier_restore(request, id):
    supplier = get_object_or_404(Supplier, id=id)
    supplier.is_active = True
    supplier.save()

    return redirect('admin_panel:supplier_list')

def genre_list(request):
    q = request.GET.get('q')

    genres = Genre.objects.all()

    if q:
        genres = genres.filter(name__icontains=q)

    return render(request, 'admin_panel//dictionaries/genres/genres.html', {
        'genres': genres
    })
def genre_create(request):
    form = GenreForm(request.POST or None)

    if form.is_valid():
        form.save()
        return redirect('admin_panel:genre_list')

    return render(request, 'admin_panel/dictionaries/genres/genre_form.html', {'form': form})

def genre_delete(request, id):
    genre = get_object_or_404(Genre, id=id)

    if genre.book_set.exists():
        genres = Genre.objects.all()

        return render(
            request,
            'admin_panel/dictionaries/genres/genres.html',
            {
                'genres': genres,
                'error': 'Нельзя удалить — используется в книгах'
            }
        )

    genre.delete()
    return redirect('admin_panel:genre_list')

def author_list(request):
    q = request.GET.get('q')

    authors = Author.objects.all()

    if q:
        authors = authors.filter(
            Q(first_name__icontains=q) |
            Q(last_name__icontains=q) |
            Q(middle_name__icontains=q)
        )

    return render(request, 'admin_panel/dictionaries/authors/authors.html', {
        'authors': authors
    })

def author_create(request):
    form = AuthorForm(request.POST or None)

    if form.is_valid():
        form.save()
        return redirect('admin_panel:author_list')

    return render(request, 'admin_panel/dictionaries/authors/author_form.html', {'form': form})

def author_delete(request, id):
    author = get_object_or_404(Author, id=id)

    if author.book_set.exists():
        authors = Author.objects.all()

        return render(
            request,
            'admin_panel/dictionaries/authors/authors.html',
            {
                'authors': authors,
                'error': 'Нельзя удалить — автор используется в книгах'
            }
        )

    author.delete()
    return redirect('admin_panel:author_list')

def tag_list(request):
    q = request.GET.get('q')

    tags = Tag.objects.all()

    if q:
        tags = tags.filter(name__icontains=q)

    return render(request, 'admin_panel/dictionaries/tags/tags.html', {
        'tags': tags
    })

def tag_create(request):
    form = TagForm(request.POST or None)

    if form.is_valid():
        form.save()
        return redirect('admin_panel:tag_list')

    return render(request, 'admin_panel/dictionaries/tags/tag_form.html', {'form': form})

def tag_delete(request, id):
    tag = get_object_or_404(Tag, id=id)

    if Book.objects.filter(tags=tag).exists():
        tags = Tag.objects.all()

        return render(
            request,
            'admin_panel/dictionaries/tags/tags.html',
            {
                'tags': tags,
                'error': 'Нельзя удалить — тег используется в книгах'
            }
        )

    tag.delete()
    return redirect('admin_panel:tag_list')

def dictionaries(request):
    return render(request, 'admin_panel/dictionaries/index.html')

@login_required
@admin_required
def product_create(request):

    product_form = ProductForm(request.POST or None)

    price_form = PriceForm(request.POST or None)
    category_slug = None
    category_form = None

    if request.method == 'POST':

        category_id = request.POST.get('category')
        category_obj = Category.objects.filter(id=category_id).first()

        category_slug = category_obj.slug if category_obj else None

        form_class = CATEGORY_MAP.get(category_slug)

        if form_class:
            category_form = form_class(request.POST)

        # validashion

        forms_valid = product_form.is_valid() and price_form.is_valid()

        if category_form:
            forms_valid = forms_valid and category_form.is_valid()

        if forms_valid:

            product = product_form.save()

            product.clear_category_data()

            price = price_form.save(commit=False)
            price.product = product
            price.save()



            # category save
            if category_form:
                obj = category_form.save(commit=False)
                obj.product = product
                obj.save()

                if hasattr(category_form, "save_m2m"):
                    category_form.save_m2m()

            # изображения
            images = request.FILES.getlist('images')
            temp_ids = request.POST.getlist('temp_ids')  # из фронта

            temp_map = {}
            new_images = []

            for i, image in enumerate(images):
                temp_id = temp_ids[i] if i < len(temp_ids) else None

                img = ProductImage.objects.create(
                    product=product,
                    image=image,
                    is_main=False
                )

                if temp_id:
                    temp_map[temp_id] = img

                new_images.append(img)

            main_image_value = request.POST.get('main_image')

            if main_image_value:
                ProductImage.objects.filter(product=product).update(is_main=False)

                if main_image_value.startswith('temp_'):
                    temp_id = main_image_value.replace('temp_', '')

                    img = temp_map.get(temp_id)
                    if img:
                        img.is_main = True
                        img.save()

            return redirect('admin_panel:product_list')

        else:
            print("PRODUCT ERRORS:", product_form.errors)
            if category_form:
                print("CATEGORY ERRORS:", category_form.errors)

    return render(request, 'admin_panel/product_form.html', {
        'product_form': product_form,
        'category_form': category_form,
        'price_form': price_form,
    })

def load_category_form(request):

    category_id = request.GET.get('category_id')
    product_id = request.GET.get('product_id')

    category = Category.objects.filter(id=category_id).first()
    if not category:
        return JsonResponse({'html': ''})

    form_class = CATEGORY_MAP.get(category.slug)
    if not form_class:
        return JsonResponse({'html': ''})

    instance = None

    if product_id:
        product = Product.objects.filter(id=product_id).first()
        if product:
            related_name = RELATED_MAP.get(category.slug)
            instance = getattr(product, related_name, None)

    form = form_class(instance=instance)

    html = render_to_string(
        'admin_panel/partials/category_form.html',
        {'category_form': form}
    )

    return JsonResponse({'html': html})

def product_list(request):
    q = request.GET.get('q', '').strip()

    active_products = Product.objects.filter(available=True)
    inactive_products = Product.objects.filter(available=False)

    if q:
        active_products = active_products.filter(
            Q(name__icontains=q) |
            Q(category__name__icontains=q)
        )

        inactive_products = inactive_products.filter(
            Q(name__icontains=q) |
            Q(category__name__icontains=q)
        )

    return render(request, 'admin_panel/product_list.html', {
        'active_products': active_products,
        'inactive_products': inactive_products,
        'q': q,
    })

def product_restore(request, id):
    product = get_object_or_404(Product, id=id)
    product.available = True
    product.save()
    return redirect('admin_panel:product_list')

# print("BOOK DATA:", request.POST.get("publisher"), type(request.POST.get("publisher")))
       # print("YEAR DATA:", request.POST.get("year"), type(request.POST.get("year")))


#
# UPDATE PRODUCT
#
 #       print(BookForm(request.POST).errors)
  #      print("POST publisher:", request.POST.get("publisher"))


@login_required
@admin_required
def product_update(request, id):

    product = get_object_or_404(Product, id=id)
    images = product.images.all()

    product_form = ProductForm(
        request.POST or None,
        instance=product
    )
    price_instance = product.prices.first()

    price_form = PriceForm(
        request.POST or None,
        instance=price_instance
    )

    # slag
    category_id = request.POST.get('category') if request.method == 'POST' else product.category_id
    category_obj = Category.objects.filter(id=category_id).first()
    category_slug = category_obj.slug if category_obj else None

    form_class = CATEGORY_MAP.get(category_slug)

    category_instance = None
    category_form = None

    if form_class:

        related_name = RELATED_MAP.get(category_slug)

        category_instance = getattr(product, related_name, None)

        category_form = form_class(
            request.POST or None,
            instance=category_instance
        )

    if request.method == 'POST':

        if product_form.is_valid() and (
            category_form.is_valid() if category_form else True
        ):

            product = product_form.save()

            product.clear_category_data()

            if price_form.is_valid():
                price_value = price_form.cleaned_data['value']

                last_price = product.prices.first()

                if not last_price or last_price.value != price_value:
                    Price.objects.create(product=product, value=price_value)
            else:
                price_value = None

            # save
            if category_form:
                obj = category_form.save(commit=False)
                obj.product = product
                obj.save()

                if hasattr(category_form, "save_m2m"):
                    category_form.save_m2m()

            # удаление изображений
            delete_ids = request.POST.getlist('delete_image')

            if delete_ids:
                ProductImage.objects.filter(
                    id__in=delete_ids,
                    product=product
                ).delete()

            # новые изображения
            uploaded_images = request.FILES.getlist('images')
            temp_ids = request.POST.getlist('temp_ids')

            temp_map = {}
            new_images = []

            for i, img_file in enumerate(uploaded_images):
                temp_id = temp_ids[i] if i < len(temp_ids) else None

                img = ProductImage.objects.create(
                    product=product,
                    image=img_file,
                    is_main=False
                )

                if temp_id:
                    temp_map[temp_id] = img

                new_images.append(img)

            # главное изображение
            main_image = request.POST.get('main_image')

            if main_image:
                product.images.update(is_main=False)

                # старые изображения
                if main_image.startswith('old_'):
                    try:
                        img_id = int(main_image.replace('old_', ''))
                        ProductImage.objects.filter(
                            id=img_id,
                            product=product
                        ).update(is_main=True)
                    except ValueError:
                        pass

                # НОВАЯ ЛОГИКА (temp_id вместо new_индексов)
                elif main_image.startswith('temp_'):
                    temp_id = main_image.replace('temp_', '')

                    img = temp_map.get(temp_id)
                    if img:
                        img.is_main = True
                        img.save()
            return redirect('admin_panel:product_list')

        else:
            print("PRODUCT FORM ERRORS:", product_form.errors)
            if category_form:
                print("CATEGORY FORM ERRORS:", category_form.errors)

    return render(request, 'admin_panel/product_form.html', {
        'product_form': product_form,
        'category_form': category_form,
        'product': product,
        'images': images,
        'price_form': price_form,
    })

def product_delete(request, id):
    product = get_object_or_404(Product, id=id)

    if product.orderitem_set.exists():
        product.available = False
        product.save()
    else:
        product.delete()

    return redirect('admin_panel:product_list')


@admin_required
def export_excel(request):

    period = request.GET.get('period', 'month')
    since = get_period(period)

    orders = Order.objects.filter(created__gte=since)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Отчет'

    ws.append(['Товар', 'Количество'])

    items = (
        OrderItem.objects
        .filter(order__in=orders)
        .values('product__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    for i in items:
        ws.append([i['product__name'], i['total']])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = f'attachment; filename=report_{period}.xlsx'

    wb.save(response)
    return response


@admin_required
def export_excel(request):
    period = request.GET.get('period', 'month')
    since = get_period(period)

    orders = Order.objects.filter(created__gte=since)

    items = OrderItem.objects.filter(
        order__in=orders
    ).select_related('product')

    wb = Workbook()
    wb.remove(wb.active)

    def write_table(ws, headers, rows):
        ws.append(headers)

        for row in rows:
            ws.append(row)

    def split_popular(qs):
        items_list = list(qs)

        avg = (
            sum(i.total for i in items_list) / len(items_list)
            if items_list else 0
        )

        popular = [i for i in items_list if i.total > avg]
        unpopular = [i for i in items_list if i.total <= avg]

        return popular, unpopular


    ws = wb.create_sheet("Рейтинг товаров")

    products_rating = Product.objects.order_by(
        '-rating',
        '-rating_count'
    )

    write_table(
        ws,
        ['Товар', 'Рейтинг', 'Отзывов'],
        [
            [p.name, p.rating, p.rating_count]
            for p in products_rating
        ]
    )

    ws = wb.create_sheet("Проданные товары")

    sold_products = (
        items.values('product__name')
        .annotate(total=Coalesce(Sum('quantity'), 0))
        .order_by('-total')
    )

    write_table(
        ws,
        ['Товар', 'Продано'],
        [
            [i['product__name'], i['total']]
            for i in sold_products
        ]
    )

    books = Book.objects.annotate(
        total=Coalesce(
            Sum(
                'product__orderitem__quantity',
                filter=Q(product__orderitem__order__in=orders)
            ),
            0
        )
    )

    popular_books, unpopular_books = split_popular(books)

    ws = wb.create_sheet("Популярные книги")

    write_table(
        ws,
        ['Книга', 'Продано'],
        [
            [b.product.name, b.total]
            for b in popular_books
        ]
    )

    ws = wb.create_sheet("Непопулярные книги")

    write_table(
        ws,
        ['Книга', 'Продано'],
        [
            [b.product.name, b.total]
            for b in unpopular_books
        ]
    )

    games = BoardGame.objects.annotate(
        total=Coalesce(
            Sum(
                'product__orderitem__quantity',
                filter=Q(product__orderitem__order__in=orders)
            ),
            0
        )
    )

    popular_games, unpopular_games = split_popular(games)

    ws = wb.create_sheet("Популярные игры")

    write_table(
        ws,
        ['Игра', 'Продано'],
        [
            [g.product.name, g.total]
            for g in popular_games
        ]
    )

    ws = wb.create_sheet("Непопулярные игры")

    write_table(
        ws,
        ['Игра', 'Продано'],
        [
            [g.product.name, g.total]
            for g in unpopular_games
        ]
    )

    stationery = Stationery.objects.annotate(
        total=Coalesce(
            Sum(
                'product__orderitem__quantity',
                filter=Q(product__orderitem__order__in=orders)
            ),
            0
        )
    )

    popular_stationery, unpopular_stationery = split_popular(
        stationery
    )

    ws = wb.create_sheet("Популярная концелярия")

    write_table(
        ws,
        ['Товар', 'Продано'],
        [
            [s.product.name, s.total]
            for s in popular_stationery
        ]
    )

    ws = wb.create_sheet("Непопулярная концелярия")

    write_table(
        ws,
        ['Товар', 'Продано'],
        [
            [s.product.name, s.total]
            for s in unpopular_stationery
        ]
    )

    authors = Author.objects.annotate(
        total=Coalesce(
            Sum(
                'book__product__orderitem__quantity',
                filter=Q(book__product__orderitem__order__in=orders)
            ),
            0
        )
    )

    popular_authors, unpopular_authors = split_popular(authors)

    ws = wb.create_sheet("Популярные авторы")

    write_table(
        ws,
        ['Автор', 'Продано'],
        [
            [str(a), a.total]
            for a in popular_authors
        ]
    )

    ws = wb.create_sheet("Непопулярные авторы")

    write_table(
        ws,
        ['Автор', 'Продано'],
        [
            [str(a), a.total]
            for a in unpopular_authors
        ]
    )

    genres = Genre.objects.annotate(
        total=Coalesce(
            Sum(
                'book__product__orderitem__quantity',
                filter=Q(book__product__orderitem__order__in=orders)
            ),
            0
        ) + Coalesce(
            Sum(
                'boardgame__product__orderitem__quantity',
                filter=Q(boardgame__product__orderitem__order__in=orders)
            ),
            0
        )
    )

    popular_genres, unpopular_genres = split_popular(genres)

    ws = wb.create_sheet("Популярные жанры")

    write_table(
        ws,
        ['Жанр', 'Продано'],
        [
            [g.name, g.total]
            for g in popular_genres
        ]
    )

    ws = wb.create_sheet("Непопулярные жанры")

    write_table(
        ws,
        ['Жанр', 'Продано'],
        [
            [g.name, g.total]
            for g in unpopular_genres
        ]
    )

    tags = Tag.objects.annotate(
        total=Coalesce(
            Sum(
                'book__product__orderitem__quantity',
                filter=Q(book__product__orderitem__order__in=orders)
            ),
            0
        )
    )

    popular_tags, unpopular_tags = split_popular(tags)

    ws = wb.create_sheet("Популярные теги")

    write_table(
        ws,
        ['Тег', 'Продано'],
        [
            [t.name, t.total]
            for t in popular_tags
        ]
    )

    ws = wb.create_sheet("Непопулярные теги")

    write_table(
        ws,
        ['Тег', 'Продано'],
        [
            [t.name, t.total]
            for t in unpopular_tags
        ]
    )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response[
        'Content-Disposition'
    ] = f'attachment; filename="report_{period}.xlsx"'

    wb.save(response)

    return response



def annotate_sales(qs, orders, relation_path):
    return qs.annotate(
        total=Coalesce(
            Sum(
                f'{relation_path}__product__orderitem__quantity',
                filter=Q(**{f'{relation_path}__product__orderitem__order__in': orders})
            ),
            0
        )
    )

def split_by_avg(qs):
    items = list(qs)
    avg = sum(i.total for i in items) / len(items) if items else 0

    return (
        [i for i in items if i.total > avg],
        [i for i in items if i.total <= avg],
    )

def write_table(ws, headers, rows):
    ws.append(headers)
    for r in rows:
        ws.append(r)

@admin_required
def dashboard(request):

    total_orders = Order.objects.count()
    total_products = Product.objects.count()

    total_sales = OrderItem.objects.aggregate(
        total=Sum('quantity')
    )['total'] or 0

    sales_day = get_sales_by_period(1)
    sales_week = get_sales_by_period(7)
    sales_month = get_sales_by_period(30)

    popular_products = (
        OrderItem.objects
        .values('product__id', 'product__name')
        .annotate(total_sold=Sum('quantity'))
        .order_by('-total_sold')[:5]
    )

    most_commented = (
        Product.objects
        .annotate(review_count=Count('reviews'))
        .order_by('-review_count')[:5]
    )

    context = {
        'total_orders': total_orders,
        'total_products': total_products,
        'total_sales': total_sales,

        'sales_day': sales_day,
        'sales_week': sales_week,
        'sales_month': sales_month,

        'popular_products': popular_products,
        'most_commented': most_commented,
    }

    return render(request, 'admin_panel/dashboard.html', context)


@admin_required
def reports(request):

    period = request.GET.get('period', 'month')
    since = get_period(period)

    orders = Order.objects.filter(
        created__gte=since
    )

    items = OrderItem.objects.filter(
        order__in=orders
    )

    def split_popular(qs):

        items_list = list(qs)

        avg = (
            sum(i.total for i in items_list) / len(items_list)
            if items_list else 0
        )

        popular = [
            i for i in items_list
            if i.total > avg
        ]

        unpopular = [
            i for i in items_list
            if i.total <= avg
        ]

        return popular, unpopular


    top_products = (
        items.values(
            'product__id',
            'product__name'
        )
        .annotate(
            total=Coalesce(
                Sum('quantity'),
                0
            )
        )
        .order_by('-total')
    )


    books = Book.objects.annotate(
        total=Coalesce(
            Sum(
                'product__orderitem__quantity',
                filter=Q(
                    product__orderitem__order__in=orders
                )
            ),
            0
        )
    )

    popular_books, unpopular_books = split_popular(
        books
    )


    games = BoardGame.objects.annotate(
        total=Coalesce(
            Sum(
                'product__orderitem__quantity',
                filter=Q(
                    product__orderitem__order__in=orders
                )
            ),
            0
        )
    )

    popular_games, unpopular_games = split_popular(
        games
    )


    stationery = Stationery.objects.annotate(
        total=Coalesce(
            Sum(
                'product__orderitem__quantity',
                filter=Q(
                    product__orderitem__order__in=orders
                )
            ),
            0
        )
    )

    popular_stationery, unpopular_stationery = split_popular(
        stationery
    )


    authors = Author.objects.annotate(
        total=Coalesce(
            Sum(
                'book__product__orderitem__quantity',
                filter=Q(
                    book__product__orderitem__order__in=orders
                )
            ),
            0
        )
    )

    popular_authors, unpopular_authors = split_popular(
        authors
    )


    genres = Genre.objects.annotate(
        total=Coalesce(
            Sum(
                'book__product__orderitem__quantity',
                filter=Q(
                    book__product__orderitem__order__in=orders
                )
            ),
            0
        ) + Coalesce(
            Sum(
                'boardgame__product__orderitem__quantity',
                filter=Q(
                    boardgame__product__orderitem__order__in=orders
                )
            ),
            0
        )
    )

    popular_genres, unpopular_genres = split_popular(
        genres
    )


    tags = Tag.objects.annotate(
        total=Coalesce(
            Sum(
                'book__product__orderitem__quantity',
                filter=Q(
                    book__product__orderitem__order__in=orders
                )
            ),
            0
        )
    )

    popular_tags, unpopular_tags = split_popular(
        tags
    )


    products_rating = Product.objects.order_by(
        '-rating',
        '-rating_count'
    )

    return render(
        request,
        'admin_panel/reports.html',
        {
            'period': period,

            'products_rating': products_rating,
            'top_products': top_products,

            'popular_books': popular_books,
            'unpopular_books': unpopular_books,

            'popular_games': popular_games,
            'unpopular_games': unpopular_games,

            'popular_stationery': popular_stationery,
            'unpopular_stationery': unpopular_stationery,

            'popular_authors': popular_authors,
            'unpopular_authors': unpopular_authors,

            'popular_genres': popular_genres,
            'unpopular_genres': unpopular_genres,

            'popular_tags': popular_tags,
            'unpopular_tags': unpopular_tags,
        }
    )